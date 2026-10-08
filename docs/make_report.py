# -*- coding: utf-8 -*-
"""Сборка отчёта по учебной практике (Плотников П.Э., ИС-943).

Запуск: python make_report.py
"""
import os
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_content import APPENDIX_FILES, BLOCKS, REFERENCES  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PRACTIC = os.path.dirname(BASE)
TPL = os.path.join(PRACTIC, "!_Стуктура отчета по УП_01.docx")
OUT = os.path.join(PRACTIC, "Отчет_Плотников_КонсалтПро.docx")
IMG = os.path.join(BASE, "report", "images")

DATE = "08.10.2026"

fig_n = 0
tab_n = 0
list_n = 0

doc = Document(TPL)

# --- очистка тела шаблона (сохраняем sectPr со стилями и полями страницы) ---
body = doc.element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)


def set_font(run, name="Times New Roman"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), name)


def ensure_code_style():
    styles = doc.styles
    try:
        st = styles["Код"]
    except KeyError:
        st = styles.add_style("Код", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
    st.font.name = "Courier New"
    st.font.size = Pt(10)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), "Courier New")
    pf = st.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.left_indent = Cm(0)
    pf.first_line_indent = Cm(0)


def add_page_break():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_footer_page_number():
    footer = doc.sections[0].footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r1._r.append(fld)
    r2 = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r2._r.append(instr)
    r3 = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = p.add_run("1")
    r5 = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)
    for r in p.runs:
        set_font(r)
        r.font.size = Pt(14)


def enable_update_fields():
    settings = doc.settings.element
    upd = settings.find(qn("w:updateFields"))
    if upd is None:
        upd = OxmlElement("w:updateFields")
        settings.append(upd)
    upd.set(qn("w:val"), "true")


def add_toc():
    p = doc.add_paragraph()
    r1 = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    fld.set(qn("w:dirty"), "true")
    r1._r.append(fld)
    r2 = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    r2._r.append(instr)
    r3 = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = p.add_run("Содержание сформируется автоматически при открытии "
                   "документа (либо выделите поле и нажмите F9).")
    r4.italic = True
    r4.font.size = Pt(12)
    r5 = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)


def add_p(text, style="УП_Осн"):
    return doc.add_paragraph(text, style=style)


def add_li(text):
    p = doc.add_paragraph(style="УП_Осн")
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    p.add_run(text)
    return p


def add_fig(fname, caption, width):
    global fig_n
    fig_n += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(os.path.join(IMG, fname), width=Cm(width))
    doc.add_paragraph(f"Рисунок {fig_n} – {caption}", style="УП_рис")


def _shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def add_table(headers, rows, caption):
    global tab_n
    tab_n += 1
    cp = doc.add_paragraph(f"Таблица {tab_n} – {caption}", style="УП_Осн")
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cp.paragraph_format.keep_with_next = True
    cp.runs[0].bold = False

    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    tblpr = t._tbl.tblPr
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:type"), "pct")
    tblw.set(qn("w:w"), "5000")
    tblpr.append(tblw)

    for j, head in enumerate(headers):
        cell = t.cell(0, j)
        cell.text = ""
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        run.font.size = Pt(11)
        set_font(run)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.line_spacing = 1.0
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        _shade(cell, "E4EFE4")

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i + 1, j)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(11)
            set_font(run)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    doc.add_paragraph(style="УП_Осн")


def add_title_line(text, size, bold, gap_after):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6 if gap_after else 0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    set_font(r)
    return p


def add_references():
    for i, (title, url) in enumerate(REFERENCES, 1):
        p = doc.add_paragraph(style="УП_Осн")
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.add_run(f"[{i}] {title} [Электронный ресурс]. "
                  f"URL: {url} (дата обращения: {DATE}) – "
                  f"Текст: электронный.")


def add_appendix(letter, name, files):
    global list_n
    add_page_break()
    p = doc.add_paragraph(f"ПРИЛОЖЕНИЕ {letter}", style="УП_Заг_1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("(обязательное)", style="УП_Осн")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.first_line_indent = Cm(0)

    for rel in files:
        full = os.path.join(BASE, rel.replace("/", os.sep))
        with open(full, encoding="utf-8") as fh:
            lines = fh.read().splitlines()
        list_n += 1
        cap = doc.add_paragraph(f"Листинг {list_n} – {rel}",
                                style="УП_Осн")
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(12)
        for line in lines:
            doc.add_paragraph(line, style="Код")
        tail = doc.add_paragraph(style="УП_Осн")
        tail.paragraph_format.space_after = Pt(6)


def main():
    ensure_code_style()
    add_footer_page_number()
    enable_update_fields()

    for kind, payload in BLOCKS:
        if kind == "title_line":
            text, size, bold, gap = payload
            add_title_line(text, size, bold, gap)
        elif kind == "title_gap":
            add_title_line("", 14 * payload, False, False)
        elif kind == "h1":
            doc.add_paragraph(payload, style="УП_Заг_1")
        elif kind == "h1_std":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(payload)
            r.bold = True
            r.font.size = Pt(16)
            set_font(r)
        elif kind == "h2":
            doc.add_paragraph(payload, style="УП_Заг_2")
        elif kind == "p":
            add_p(payload)
        elif kind == "li":
            add_li(payload)
        elif kind == "fig":
            add_fig(*payload)
        elif kind == "table":
            add_table(*payload)
        elif kind == "pb":
            add_page_break()
        elif kind == "toc":
            add_toc()
        else:
            raise ValueError(f"unknown block: {kind}")

    add_references()

    for letter, name, files in APPENDIX_FILES:
        add_appendix(letter, name, files)

    doc.save(OUT)
    print("saved:", OUT)
    print(f"figures: {fig_n}, tables: {tab_n}, listings: {list_n}")


if __name__ == "__main__":
    main()
