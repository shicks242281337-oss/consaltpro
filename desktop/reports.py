"""Формирование отчётов в файлах Excel и TXT."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side


def _cell_style(cell, header=False):
    cell.alignment = Alignment(
        horizontal="center" if header else "left", vertical="center", wrap_text=True
    )
    thin = Side(style="thin", color="999999")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if header:
        cell.font = Font(bold=True, size=11)


def write_excel(path, caption, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"[:31]
    ws.append([caption])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = ws.cell(row=1, column=1)
    title.font = Font(bold=True, size=13)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        _cell_style(cell, header=True)
    for row in rows:
        ws.append(list(row))
    for row in ws.iter_rows(min_row=4, max_row=3 + len(rows), max_col=len(headers)):
        for cell in row:
            _cell_style(cell)
    for idx, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows] or [10])
        ws.column_dimensions[ws.cell(row=3, column=idx).column_letter].width = min(
            max(width + 2, 12), 55
        )
    wb.save(path)


def write_txt(path, caption, headers, rows):
    def norm(value):
        if isinstance(value, float):
            return f"{value:,.2f}".replace(",", " ")
        return str(value)

    table = [list(headers)] + [[norm(v) for v in row] for row in rows]
    widths = [max(len(row[i]) for row in table) for i in range(len(headers))]

    def line(values):
        return " | ".join(
            str(v).ljust(widths[i]) for i, v in enumerate(values)
        )

    separator = "-+-".join("-" * w for w in widths)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(caption + "\n")
        fh.write("=" * len(caption) + "\n\n")
        fh.write(line(headers) + "\n")
        fh.write(separator + "\n")
        for row in table[1:]:
            fh.write(line(row) + "\n")
        fh.write(separator + "\n")
        fh.write(f"Всего строк: {len(rows)}\n")
